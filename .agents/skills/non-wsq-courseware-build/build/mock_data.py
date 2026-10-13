"""Deterministic synthetic maintenance inputs, templates and rejoin checkpoints."""
from pathlib import Path
from datetime import datetime,timedelta
from collections import Counter
import csv,json
from openpyxl import Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,Reference
from docx import Document
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
REPORT=datetime(2026,9,30)
SITES=['Woodlands Hub','Jurong Depot','Tampines Centre','Paya Lebar Office','Changi Workshop']
CATEGORIES=['HVAC','Electrical','Plumbing','Fire Safety','General Repair']
STATUSES=['New','Assigned','In Progress','On Hold','Closed','Cancelled']
TECHS=['Aisha Tan','Ben Lim','Chloe Goh','Daniel Lee','Elena Ong','Farid Rahman','Grace Ng','Hari Kumar']
HEADERS=['Work_Order_ID','Site','Job_Category','Priority','Status','Assigned_To','Opened_Date','Due_Date','Closed_Date','Estimated_Hours','Actual_Hours','Cost_SGD','Description']
ROWS=[]
for i in range(120):
 status='Closed' if i%5 in (0,1) else ('Cancelled' if i%10==9 else ['New','Assigned','In Progress','On Hold'][i%4])
 opened=REPORT-timedelta(days=5+i%80);due=opened+timedelta(days=7+i%10)
 closed=due+timedelta(days=(i//5)%5-2) if status=='Closed' else None
 # Keep every closure at or before the reporting date.
 if closed and closed>REPORT:closed=REPORT
 owner='' if status=='New' and i%3==0 else TECHS[i%8]
 ROWS.append([f'WO-2026-{i+1:04d}',SITES[(i//3)%5],CATEGORIES[(i//4)%5],['Low','Medium','High','Critical'][i%4],status,owner,opened,due,closed,2+i%7,(1+i%9) if status=='Closed' else 0,round(90+i%13*35.5,2),f'{CATEGORIES[(i//4)%5]} scheduled or reactive task {i+1:03d}'])
def metrics(rows=ROWS):
 closed=[r for r in rows if r[4]=='Closed'];back=[r for r in rows if r[4] not in ('Closed','Cancelled')]
 return {'total':len(rows),'closed':len(closed),'cancelled':sum(r[4]=='Cancelled' for r in rows),'backlog':len(back),'overdue':sum(r[7]<REPORT for r in back),'unassigned_backlog':sum(not r[5] for r in back),'on_time_closed':sum(r[8]<=r[7] for r in closed),'on_time_rate':sum(r[8]<=r[7] for r in closed)/len(closed),'cost':round(sum(r[11] for r in rows),2),'site_counts':dict(Counter(r[1] for r in rows)),'category_counts':dict(Counter(r[2] for r in rows))}
EXPECTED=metrics()
def style(ws):
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='1F6FEB');c.font=Font(color='FFFFFF',bold=True)
 for col in ws.columns:
  ws.column_dimensions[col[0].column_letter].width=min(45,max(17,len(str(col[0].value))+2))
 for row in ws:
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(c.value,datetime):c.number_format='yyyy-mm-dd'
def table(ws,name):
 t=Table(displayName=name,ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
def add_dashboard(wb,ready=True):
 ws=wb.create_sheet('Dashboard');ws.append(['Northstar Maintenance Dashboard','Selection','As at','2026-09-30'])
 for label in ['Site','Job_Category','Status','Assigned_To']:ws.append([label,'All'])
 ws.append(['Measure','Current selection']); ws.append(['Total work orders']);ws.append(['Backlog']);ws.append(['Overdue backlog']);ws.append(['Closed']);ws.append(['On-time closure rate']);ws.append(['Cost SGD'])
 # SUMPRODUCT plus structured references includes rows appended to the table.
 mask='(IF($B$2="All",1,--(tblWorkOrders[Site]=$B$2)))*(IF($B$3="All",1,--(tblWorkOrders[Job_Category]=$B$3)))*(IF($B$4="All",1,--(tblWorkOrders[Status]=$B$4)))*(IF($B$5="All",1,--(tblWorkOrders[Assigned_To]=$B$5)))'
 formula={7:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Work_Order_ID]<>""))',8:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]<>"Closed")*--(tblWorkOrders[Status]<>"Cancelled"))',9:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]<>"Closed")*--(tblWorkOrders[Status]<>"Cancelled")*--(tblWorkOrders[Due_Date]<Control!$B$2))',10:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]="Closed"))',11:f'=IF(B10=0,"N/A",SUMPRODUCT({mask}*--(tblWorkOrders[Status]="Closed")*--(tblWorkOrders[Closed_Date]<=tblWorkOrders[Due_Date]))/B10)',12:f'=SUMPRODUCT({mask}*tblWorkOrders[Cost_SGD])'}
 lists=wb.create_sheet('Filter_Lists')
 for col,values in enumerate([SITES,CATEGORIES,STATUSES,TECHS],1):
  for row,value in enumerate(['All']+values,1):lists.cell(row,col,value)
  from openpyxl.utils import get_column_letter
  letter=get_column_letter(col); dv=DataValidation(type='list',formula1=f"'Filter_Lists'!${letter}$1:${letter}${len(values)+1}");ws.add_data_validation(dv);dv.add(ws.cell(col+1,2))
 if ready:
  for row,f in formula.items():ws.cell(row,2,f)
  ws['B11'].number_format='0.0%';ws['B12'].number_format='"S$"#,##0.00'
  for start,title,col,vals in [(15,'Work orders by status','Status',STATUSES),(24,'Work orders by category','Job_Category',CATEGORIES),(33,'Work orders by site','Site',SITES)]:
   ws.cell(start,1,title);ws.cell(start,2,'Count')
   for k,value in enumerate(vals,start+1):ws.cell(k,1,value);ws.cell(k,2,f'=SUMPRODUCT({mask}*--(tblWorkOrders[{col}]=A{k}))')
   ch=BarChart();ch.title=title;ch.add_data(Reference(ws,min_col=2,min_row=start,max_row=start+len(vals)),titles_from_data=True);ch.set_categories(Reference(ws,min_col=1,min_row=start+1,max_row=start+len(vals)));ch.width=21;ch.height=8;ws.add_chart(ch,f'E{start-13}')
 style(ws);ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=24
 lists.sheet_state='hidden'
def workbook(path,raw=False,dashboard=False,blank=False):
 wb=Workbook();ws=wb.active;ws.title='About_This_Data';ws.append(['Field','Value']);ws.append(['Company','Northstar Facilities - fictional']);ws.append(['Reporting date',REPORT]);ws.append(['Scope','120 Q3 2026 work orders, five sites. All identities and records are synthetic.']);ws.append(['Meaning','Cost is cumulative recorded cost of all statuses; not a budget or productivity claim.']);ws.append(['Rejoin','This input is a deterministic checkpoint, not proof of your Copilot output.'])
 ws=wb.create_sheet('Work_Orders');ws.append(HEADERS)
 source=[list(r) for r in ROWS]
 if raw:
  source[3][2]=' hvac ';source[12][2]='Fire safety ';source += [list(source[4]),list(source[19])]
 for r in source:ws.append(r)
 table(ws,'tblWorkOrders')
 ws=wb.create_sheet('Technicians');ws.append(['Technician','Primary_Category','Available_Hours_Week','Home_Site'])
 for i,t in enumerate(TECHS):ws.append([t,CATEGORIES[(i//4)%5],30 if i%2 else 35,SITES[i%5]])
 table(ws,'tblTechnicians')
 ws=wb.create_sheet('Control');ws.append(['Metric','Expected / control']);ws.append(['Report_Date',REPORT])
 for k,v in EXPECTED.items():
  if not isinstance(v,dict):ws.append([k,v])
 ws.append(['Backlog definition','Status is neither Closed nor Cancelled']);ws.append(['Overdue definition','Backlog and Due_Date strictly earlier than Report_Date']);ws.append(['On-time definition','Closed_Date <= Due_Date; denominator Closed only']);ws.append(['Unknown assignment','Blank Assigned_To stays blank; propose separately'])
 ws=wb.create_sheet('Prompt_Log');ws.append(['Step','App / mode','Prompt','What Copilot did','How I verified','Decision'])
 ws=wb.create_sheet('Checks');ws.append(['Check','Formula / evidence','Result']);ws.append(['Total clean rows','=ROWS(tblWorkOrders[Work_Order_ID])',''])
 for ws in wb:style(ws)
 if dashboard:add_dashboard(wb,ready=not blank)
 path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
def doc(path,title,paras):
 d=Document();d.add_heading(title,0)
 for heading,body in paras:d.add_heading(heading,1);d.add_paragraph(body)
 d.save(path)
def deck(path,checkpoint=False,count=8):
 from pptx.chart.data import CategoryChartData
 from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION
 r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
 titles=['Q3 maintenance: prioritise the overdue backlog','120 orders: reconcile the portfolio','Site backlog: compare open work','Status and category mix','30 of 48 closures were on time','6 backlog orders have no current owner','Proposed recovery actions and trade-offs','Decision requested: review the recovery proposal']
 if count==4:titles=['Northstar maintenance portfolio - 30 September','50 of 60 backlog orders are overdue','6 backlog orders need an assignment proposal','Decision requested: review the proposed priorities']
 bodies=[
 'Reporting date:30 September 2026. All sites/categories/statuses/owners. Synthetic training snapshot. Findings describe recorded work; actions below are proposed.',
 f"Closed {EXPECTED['closed']} + Cancelled {EXPECTED['cancelled']} + Backlog {EXPECTED['backlog']} = {EXPECTED['total']}. Recorded cost S${EXPECTED['cost']:,.2f}; this is not a budget or saving.",
 'Backlog excludes Closed and Cancelled. Compare counts across five sites before choosing recovery priorities. Frequency alone does not establish the cause.',
 'Status describes current workflow state; category describes job type. All-category counts sum to 120; category frequency is not a root-cause diagnosis.',
 f"On-time closure rate = {EXPECTED['on_time_closed']}/{EXPECTED['closed']} = {EXPECTED['on_time_rate']:.1%}. Overdue backlog: {EXPECTED['overdue']}, using Due_Date < 30 September 2026. Closed-only denominator, not all 120 orders.",
 f"Unassigned backlog: {EXPECTED['unassigned_backlog']}. Compare primary category and available hours across 8 synthetic technicians. Current owners remain unchanged; specific suitable/approved allocations UNKNOWN.",
 'Proposed: review the overdue queue by priority/site; validate constraints before suggesting owners; keep category/capacity evidence beside each proposal. Trade-off: backlog recovery may compete with newly opened work. Savings and approvals UNKNOWN.',
 'Requested decision: review and approve or revise the proposed recovery priorities and assignment checks. Decision owner, approval outcome and next review date UNKNOWN until supplied. No email sent and no source assignment changed.'
 ]
 if count==4:bodies=[bodies[1],bodies[4],bodies[5],bodies[7]]
 notes=[
 'Source tblWorkOrders, all filters All, as at 2026-09-30. State scope before interpretation. This is a deterministic rejoin example, not evidence of Copilot execution.',
 'Counts partition the selected source table exactly; recorded cost includes all statuses. Do not infer a budget or savings.',
 'Backlog chart counts Status neither Closed nor Cancelled. Site totals are source-derived and reconcile to 60. Counts alone do not explain causes.',
 'Status chart uses all 120 source rows; category comparison is nominal. Check units and avoid claims of causation.',
 'Closed on/before Due_Date is on time. Numerator 30 and denominator 48 are both source-derived. A different scope must be recalculated.',
 'Six blank current owners are in backlog. Technicians table contains planning inputs, not approval or comprehensive safety/capacity evidence.',
 'Actions are proposals linked to observed overdue/unassigned work. They need operational review; no promised saving or productivity gain.',
 'Ask for a decision without claiming one was made. The PowerPoint is a static snapshot; retain the interactive workbook and evidence register.'
 ]
 if count==4:notes=[notes[1],notes[4],notes[5],notes[7]]
 for i,title in enumerate(titles):
  sh=r.slides.add_slide(r.slide_layouts[6]);tb=sh.shapes.add_textbox(Inches(.7),Inches(.5),Inches(12),Inches(1));para=tb.text_frame.paragraphs[0];para.text=title;para.font.size=Pt(30);para.font.bold=True;para.font.color.rgb=RGBColor.from_string('0B3060')
  tb=sh.shapes.add_textbox(Inches(.8),Inches(1.65),Inches(11.7),Inches(1.5 if checkpoint else 3.6));tb.text_frame.word_wrap=True;para=tb.text_frame.paragraphs[0];para.text=bodies[i] if checkpoint else 'Add the verified chart or evidence table for this slide. Separate observed facts from proposed actions.';para.font.size=Pt(20)
  if checkpoint and ((count==8 and i in (2,3)) or (count==4 and i==1)):
   if count==8 and i==3:counts=Counter(row[4] for row in ROWS);label='All work orders by status'
   else:counts=Counter(row[1] for row in ROWS if row[4] not in ('Closed','Cancelled'));label='Backlog work orders by site'
   data=CategoryChartData();data.categories=list(counts);data.add_series(label,list(counts.values()));chart=sh.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(1),Inches(3.3),Inches(11),Inches(3.2),data).chart;chart.has_legend=False;chart.plots[0].has_data_labels=True;chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;chart.category_axis.tick_labels.font.size=Pt(12);chart.value_axis.tick_labels.font.size=Pt(12);chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string('13A89E')
  elif checkpoint:
   tb=sh.shapes.add_textbox(Inches(.8),Inches(3.6),Inches(11.7),Inches(2));para=tb.text_frame.paragraphs[0];para.text=['Observation -> evidence -> proposed action -> decision','Keep status, site, category and owner scope consistent.','Validate constraints before confirming an operational change.','Static presentation: verified snapshot, not an automatic live link.'][i%4];para.font.size=Pt(27);para.font.bold=True;para.font.color.rgb=RGBColor.from_string('1F6FEB')
  tb=sh.shapes.add_textbox(Inches(.8),Inches(6.8),Inches(11.7),Inches(.35));para=tb.text_frame.paragraphs[0];para.text=f'C27 v2.0 | Synthetic Northstar example | 2026-09-30 | All filters All | Slide{i+1}';para.font.size=Pt(10)
  sh.notes_slide.notes_text_frame.text=notes[i] if checkpoint else 'Provide the selected verified snapshot and replace scaffold content. Add source/date/scope notes; do not invent facts.'
 r.save(path)
def build_lab_data(num,folder):
 root=Path(folder);data=root/'data';tmp=root/'templates';cp=root/'checkpoints'
 for p in [data,tmp,cp]:p.mkdir(parents=True,exist_ok=True)
 workbook(data/'Northstar_Work_Orders.xlsx',raw=(num==1),dashboard=(num>=6))
 workbook(cp/'Northstar_Verified_Checkpoint.xlsx',dashboard=(num>=5))
 with (data/'Northstar_Work_Orders.csv').open('w',newline='') as f:
  w=csv.writer(f);w.writerow(HEADERS);w.writerows([[v.strftime('%Y-%m-%d') if isinstance(v,datetime) else v for v in r] for r in ROWS])
 (data/'verified-controls.json').write_text(json.dumps(EXPECTED,indent=2)+'\n')
 (data/'Northstar_Maintenance_Context.txt').write_text('Northstar Facilities is fictional. Reporting date: 2026-09-30. Work orders have statuses New, Assigned, In Progress, On Hold, Closed, Cancelled. Backlog excludes Closed and Cancelled. Overdue means backlog due before the reporting date. Unassigned jobs require a proposal, never an invented confirmed owner. Proposed actions need a manager decision. All sample data is synthetic.\n')
 with (tmp/'Evidence_Register.csv').open('w',newline='') as f:csv.writer(f).writerows([['Claim','Value','Source_File','Sheet_Table','Filter_Scope','As_At','Verified_By','Proposed_or_Observed']])
 with (tmp/'Assignment_Proposals.csv').open('w',newline='') as f:csv.writer(f).writerows([['Work_Order_ID','Current_Assigned_To','Proposed_Assigned_To','Category_Match','Capacity_Check','Reason','Manager_Decision']])
 doc(tmp/'Northstar_Operations_Brief_Template.docx','Northstar operations brief',[('Reporting scope','Record file, table, filters and reporting date.'),('Verified measures','Replace this paragraph with verified values and denominators.'),('Backlog and priority','Summarise evidence without guessing causal explanations.'),('Assignment proposals','Keep current assignment and proposed assignment separate.'),('Decision and open questions','Name the decision needed; use UNKNOWN for missing facts.')])
 doc(cp/'Northstar_Verified_Snapshot.docx','Northstar verified maintenance snapshot',[('Source and scope','Synthetic tblWorkOrders: 120 records; all filters All; 30 September 2026.'),('Verified controls',json.dumps(EXPECTED,indent=2)),('Interpretation','Values describe this fictional snapshot. Category mix is a count, not evidence of root cause. Proposals are unapproved.'),('Presentation outline','Portfolio; backlog; mix; SLA; capacity; actions; decision. Reconcile every figure against the dashboard before use.')])
 doc(tmp/'Northstar_Presentation_Brief_Template.docx','Northstar management presentation brief',[('Audience and decision','Audience: facilities management. Decision: prioritise a proposed backlog recovery plan; approvals UNKNOWN.'),('Slide evidence','For each of eight slides specify headline, figure, chart, scope and source.'),('Limitations','Snapshot date, synthetic data, N/A for zero denominators, no invented savings.'),('Narrative','Observation -> operational implication -> proposed action -> decision request.')])
 deck(tmp/'Northstar_Management_Deck_Template.pptx')
 if num==4:deck(cp/'Northstar_Initial_Update_Checkpoint.pptx',True,count=4)
 if num>=7:deck(cp/'Northstar_Management_Example.pptx',True)
 (tmp/'Northstar_Outlook_Draft_Template.txt').write_text('To: operations.manager@example.invalid\nSubject: DRAFT - Northstar maintenance update - review required\nScope/date:\nVerified findings:\nProposed assignment/actions (unapproved):\nDecision needed:\nAttachment/source links:\nDo not send this training draft.\n')
 (data/'Northstar_Email_Thread.txt').write_text('TRAINING THREAD - synthetic; do not send.\n2026-09-29 09:00 Operations manager: Please prepare the 30 September portfolio snapshot; show open work by site, category and current assignment.\n2026-09-29 10:15 Supervisor: Check overdue jobs before proposing reassignment. Available hours are planning inputs, not approval.\n2026-09-30 08:30 Analyst: The source export is in the workbook. No new assignment has been approved. Please produce a Word brief and a short PowerPoint update, then prepare a draft reply.\n')
print('Northstar controls:',EXPECTED) if __name__=='__main__' else None
